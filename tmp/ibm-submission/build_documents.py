from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.platypus import Paragraph, Table, TableStyle


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "output" / "ibm-submission"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = "14234B"
BLUE = "0F62FE"
PALE = "EEF4FF"
TEXT = "202938"
LIGHT = "D9E2F2"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def build_concept_note():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].font.color.rgb = RGBColor.from_string(TEXT)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(30)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    title_style_pr = styles["Title"].element.get_or_add_pPr()
    inherited_border = title_style_pr.find(qn("w:pBdr"))
    if inherited_border is not None:
        title_style_pr.remove(inherited_border)
    for name, size in (("Heading 1", 17), ("Heading 2", 12)):
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.add_run("Mentora AI Concept Note")
    title_pr = title._p.get_or_add_pPr()
    title_border = title_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_pr.remove(title_border)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = subtitle.add_run("IBM SkillsBuild for Academia Internship Project")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    meta = doc.add_paragraph("SDG 4 Quality Education  |  Agentic AI Learning Platform  |  October 2026")
    meta.paragraph_format.space_after = Pt(14)

    intro = doc.add_paragraph()
    r = intro.add_run("Project summary. ")
    r.bold = True
    intro.add_run(
        "Mentora AI is a working web prototype that personalizes learning around a student's career goal, prior skills, assessment results, and ongoing performance. The system combines specialist AI agents, grounded retrieval, adaptive quizzes, career guidance, project recommendations, reminders, and progress tracking in one learner experience."
    )

    add_heading(doc, "Problem and opportunity")
    doc.add_paragraph(
        "Students often move between videos, course platforms, search engines, and general chat tools without a coherent learning plan. Static courses rarely adjust to career goals or changing performance. New AI developments also make learning material age quickly, while limited mentoring and weak accountability reduce completion. These gaps directly affect SDG 4 because access alone does not guarantee relevant, effective, and inclusive learning."
    )
    add_bullets(doc, [
        "Learning resources remain scattered across unrelated platforms.",
        "Students receive limited feedback on weak topics and next steps.",
        "Career advice, portfolio projects, and learning progress are usually disconnected.",
        "Generic chat responses can lack grounding, learner context, and persistent evidence.",
    ])

    add_heading(doc, "Proposed solution")
    doc.add_paragraph(
        "Mentora AI provides a single learner workspace. A student creates a profile, chooses a career goal, completes an initial assessment, and receives a sequenced roadmap. The platform records progress and quiz evidence, routes requests to specialist agents, retrieves relevant material from a curated knowledge base, and recommends skill gaps and portfolio projects."
    )
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    widths = [Inches(1.55), Inches(2.3), Inches(2.95)]
    headers = ["Capability", "Learner value", "Current prototype"]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        run = p.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
    rows = [
        ("Personalization", "A plan shaped by goals and evidence", "Profile, assessment bands, adaptive roadmap"),
        ("Teaching", "Contextual explanations with traceable sources", "Tutor routing with local ChromaDB retrieval"),
        ("Assessment", "Difficulty changes with performance", "Persisted quizzes, explanations, weak-topic updates"),
        ("Career support", "Clear skills and portfolio direction", "Skill gaps, certifications, interview preparation, ranked projects"),
        ("Engagement", "Visible consistency and follow-up", "Activity history, streaks, reminders, dashboard"),
    ]
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].width = widths[i]
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            if row_index % 2 == 0:
                set_cell_shading(cells[i], "F5F8FD")
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.add_run(value)
    set_table_borders(table)

    add_heading(doc, "Agentic system design")
    doc.add_paragraph(
        "The orchestrator classifies each request and selects a specialist. The Tutor explains concepts using retrieved knowledge. Quiz Master directs students to scored adaptive quizzes. Career Coach uses the learner profile to explain skill gaps. Project Guide recommends feasible portfolio work. Researcher presents curated AI updates. Reminder Agent helps the learner create an actionable follow-up."
    )
    add_bullets(doc, [
        "React and TypeScript provide the responsive learner interface.",
        "FastAPI exposes authenticated APIs and orchestrates domain services.",
        "SQLAlchemy and Alembic manage learner records, assessments, roadmaps, quizzes, projects, activities, and reminders.",
        "ChromaDB stores local knowledge chunks with deterministic offline embeddings.",
        "A provider boundary supports a no-key deterministic demo and optional OpenAI responses.",
    ])

    add_heading(doc, "Implemented learner journey")
    steps = [
        ("1", "Profile", "Record education, experience, skills, career goal, and weekly study time."),
        ("2", "Assessment", "Measure topic performance and classify strong, average, and weak areas."),
        ("3", "Roadmap", "Create an ordered path with locked progression and backend-owned completion."),
        ("4", "Learn and practise", "Ask a grounded mentor and complete adaptive, scored quizzes."),
        ("5", "Career evidence", "Review skill gaps, projects, certifications, progress, and activity streaks."),
    ]
    st = doc.add_table(rows=1, cols=3)
    st.autofit = False
    for i, (label, width) in enumerate(zip(["Step", "Stage", "Outcome"], [0.55, 1.25, 5.0])):
        cell = st.rows[0].cells[i]
        cell.width = Inches(width)
        set_cell_shading(cell, BLUE)
        rr = cell.paragraphs[0].add_run(label)
        rr.bold = True
        rr.font.color.rgb = RGBColor(255, 255, 255)
    for n, stage, outcome in steps:
        cells = st.add_row().cells
        for i, value in enumerate((n, stage, outcome)):
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[i].paragraphs[0].add_run(value)
    for column_index, width in enumerate((0.55, 1.25, 5.0)):
        for cell in st.columns[column_index].cells:
            cell.width = Inches(width)
    set_table_borders(st)

    add_heading(doc, "Target users and adoption")
    doc.add_paragraph(
        "Primary users are college and engineering students, including learners entering AI from non-technical backgrounds. Early adopters include AI clubs, IBM SkillsBuild learners, colleges with AI labs, and students preparing for AI careers. Institutions, CSR education programs, and government skill initiatives can support broader adoption after a validated pilot."
    )

    add_heading(doc, "Validation and evidence")
    doc.add_paragraph(
        "The current submission includes an operational local MVP rather than a concept-only mock-up. Automated verification covers authentication, user isolation, assessment and roadmap personalization, adaptive quizzes, retrieval, agent routing, career analysis, projects, reminders, activity, and streaks."
    )
    metrics = doc.add_table(rows=1, cols=2)
    metrics.autofit = False
    for i, text in enumerate(("Verification", "Result")):
        cell = metrics.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        rr = cell.paragraphs[0].add_run(text)
        rr.bold = True
        rr.font.color.rgb = RGBColor(255, 255, 255)
    for left, right in [
        ("Backend automated tests", "12 passed"),
        ("Frontend unit tests", "3 passed"),
        ("Frontend quality checks", "Lint, TypeScript, and production build passed"),
        ("Runtime smoke test", "Authenticated core routes and grounded mentor response passed"),
    ]:
        cells = metrics.add_row().cells
        cells[0].paragraphs[0].add_run(left)
        cells[1].paragraphs[0].add_run(right)
        for cell in cells:
            set_cell_margins(cell)
    set_table_borders(metrics)

    add_heading(doc, "Impact and success measures")
    doc.add_paragraph(
        "A pilot should evaluate whether Mentora helps learners complete planned topics, improve quiz performance, return consistently, and produce career evidence. The first evaluation should track roadmap completion, quiz improvement by topic, active learning days, project progress, and learner feedback. Institutional adoption and reminder engagement become relevant only after those channels exist."
    )

    add_heading(doc, "Responsible scope and limitations")
    doc.add_paragraph(
        "The prototype keeps offline behavior honest. It uses curated Markdown knowledge, local retrieval, deterministic personalization, and in-app reminders. Research updates are curated records. Live OpenAI mode requires a configured key and separate quality evaluation. WhatsApp, email, automated crawling, native mobile, LMS, payments, and marketplaces remain future integrations. Mentora AI is an internship project and is not presented as an official IBM product."
    )

    add_heading(doc, "Implementation plan")
    add_bullets(doc, [
        "Pilot the current web MVP with a small student cohort and collect structured feedback.",
        "Measure retrieval quality and learning outcomes before expanding the knowledge corpus.",
        "Add monitored live-model evaluation, production authentication hardening, PostgreSQL, backups, and shared rate limiting before public deployment.",
        "Assess consent, privacy, and message-delivery requirements before enabling email or WhatsApp reminders.",
    ])

    add_heading(doc, "Conclusion")
    doc.add_paragraph(
        "Mentora AI turns the IBM SkillsBuild example into a tested prototype centered on learner evidence. The product demonstrates how agent routing, retrieval, adaptive assessment, and career guidance can work together while keeping future commercial integrations clearly separated from the current MVP."
    )

    for section in doc.sections:
        footer = section.footer
        footer.paragraphs[0].text = ""
        ft = footer.add_table(rows=1, cols=2, width=Inches(6.9))
        ft.autofit = False
        ft.columns[0].width = Inches(6.25)
        ft.columns[1].width = Inches(0.65)
        left = ft.cell(0, 0).paragraphs[0]
        left.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = left.add_run("Mentora AI  |  IBM SkillsBuild internship project")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 110, 125)
        right = ft.cell(0, 1).paragraphs[0]
        add_page_number(right)
        if right.runs:
            right.runs[0].font.size = Pt(8)
            right.runs[0].font.color.rgb = RGBColor(100, 110, 125)

    path = OUT / "Mentora_AI_Concept_Note.docx"
    doc.save(path)
    return path


def build_lean_canvas():
    width, height = landscape(A4)
    path = OUT / "Mentora_AI_Lean_Canvas.pdf"
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=(width, height))
    c.setFillColor(colors.HexColor("#F7F9FE"))
    c.rect(0, 0, width, height, fill=1, stroke=0)
    c.setFillColor(colors.HexColor("#14234B"))
    c.rect(0, height - 25 * mm, width, 25 * mm, fill=1, stroke=0)
    c.setFont("Helvetica-Bold", 24)
    c.setFillColor(colors.white)
    c.drawString(14 * mm, height - 15.5 * mm, "Mentora AI Lean Canvas")
    c.setFont("Helvetica", 9)
    c.drawRightString(width - 14 * mm, height - 15 * mm, "IBM SkillsBuild internship project | SDG 4 Quality Education")

    styles = {
        "head": ParagraphStyle("head", fontName="Helvetica-Bold", fontSize=10.3, leading=12, textColor=colors.HexColor("#0F62FE"), spaceAfter=4),
        "body": ParagraphStyle("body", fontName="Helvetica", fontSize=7.7, leading=9.4, textColor=colors.HexColor("#202938"), alignment=TA_LEFT),
    }

    def block(title, items):
        bullets = "<br/>".join(f"&#8226; {item}" for item in items)
        return [Paragraph(title, styles["head"]), Paragraph(bullets, styles["body"])]

    top = [
        block("1 Problem", ["Learning resources are scattered", "Static courses do not adapt to goals or performance", "Students lack consistent mentoring and accountability", "AI content becomes outdated quickly"]),
        block("4 Solution", ["Personalized roadmaps", "Grounded AI tutor", "Adaptive assessments", "Career and project guidance", "Progress, streaks, and reminders"]),
        block("3 Unique value", ["One learner record connects assessment, teaching, practice, career preparation, and progress", "Specialist agents respond using persistent learner context", "Offline demo remains usable without external keys"]),
        block("9 Unfair advantage", ["Working end-to-end MVP", "Transparent agent routing and grounded sources", "Career-aware personalization", "SDG 4 alignment for institutional pilots"]),
        block("2 Customer segments", ["College and engineering students", "Non-technical AI learners", "Universities and AI clubs", "CSR and government skill programs"]),
    ]
    lower = [
        block("5 Channels", ["College and university pilots", "IBM SkillsBuild communities", "Student ambassadors", "Web application", "Future LMS, mobile, email, and WhatsApp"]),
        block("8 Key metrics", ["Roadmap completion", "Quiz improvement by topic", "Active learning days and streaks", "Project progress", "Learner satisfaction and retention"]),
        block("7 Cost structure", ["Model and embedding usage", "Database and hosting", "Knowledge validation", "Product development", "Support and institutional adoption"]),
        block("6 Revenue streams", ["Freemium learner plan", "Premium mentor features", "College licenses", "CSR and government programs", "Future certification and marketplace fees"]),
    ]

    margin = 12 * mm
    usable_w = width - 2 * margin
    top_y = height - 31 * mm
    top_h = 100 * mm
    bottom_h = 55 * mm
    top_table = Table([top], colWidths=[usable_w / 5] * 5, rowHeights=[top_h])
    top_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#B7C6E2")),
        ("INNERGRID", (0, 0), (-1, -1), 0.8, colors.HexColor("#B7C6E2")),
        ("BACKGROUND", (0, 0), (-1, -1), colors.white),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    top_table.wrapOn(c, usable_w, top_h)
    top_table.drawOn(c, margin, top_y - top_h)

    lower_table = Table([lower], colWidths=[usable_w / 4] * 4, rowHeights=[bottom_h])
    lower_table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#B7C6E2")),
        ("INNERGRID", (0, 0), (-1, -1), 0.8, colors.HexColor("#B7C6E2")),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EEF4FF")),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    lower_y = top_y - top_h - 4 * mm
    lower_table.wrapOn(c, usable_w, bottom_h)
    lower_table.drawOn(c, margin, lower_y - bottom_h)

    c.setFont("Helvetica", 7.5)
    c.setFillColor(colors.HexColor("#5D6878"))
    c.drawString(margin, 7 * mm, "Current MVP: web app, deterministic personalization, local ChromaDB retrieval, in-app reminders, curated research updates")
    c.drawRightString(width - margin, 7 * mm, "Future scope is labeled explicitly")
    c.save()
    return path


if __name__ == "__main__":
    print(build_concept_note())
    print(build_lean_canvas())
